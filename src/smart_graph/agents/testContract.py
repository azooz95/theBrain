import os
import logging
import pandas as pd
from uuid import uuid4
from docx import Document
from transformers import pipeline
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

# Configure logging
logging.basicConfig(
    filename="contract_generator.log",
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s"
)

# Session state (mocked; use DB or Redis in production)
CONTRACT_SESSION = {}

class ContractGenerator:
    def __init__(self, template_path: str, contract_type: str = "Generic"):
        self.template_path = template_path
        self.contract_type = contract_type
        self.placeholders = []
        self.responses = {}

        os.makedirs("outputs", exist_ok=True)
        self.output_base = os.path.join("outputs", f"contract_output_{uuid4().hex[:4]}")
        self.docx_output = f"{self.output_base}.docx"
        self.pdf_output = f"{self.output_base}.pdf"

        try:
            self.question_generator = pipeline("text2text-generation", model="google/flan-t5-base")
        except Exception as e:
            logging.error(f"Error loading Hugging Face model: {e}")
            raise RuntimeError("Model loading failed.")

    def extract_placeholders(self) -> list:
        try:
            document = Document(self.template_path)
            for para in document.paragraphs:
                if "[" in para.text and "]" in para.text:
                    field = para.text[para.text.find("[") + 1:para.text.find("]")]
                    self.placeholders.append(field.strip())
            self.placeholders = list(set(self.placeholders))  # remove duplicates
            return self.placeholders
        except Exception as e:
            logging.error(f"Failed to extract placeholders: {e}")
            raise

    def generate_questions(self) -> dict:
        questions = {}
        for placeholder in self.placeholders:
            contextual = self._generate_contextual_question(placeholder)
            if contextual == "FALLBACK":
                prompt = f"Generate a professional question for a contract about '{placeholder}'."
                response = self.question_generator(prompt)[0]['generated_text']
                questions[placeholder] = self._postprocess_question(response)
            else:
                questions[placeholder] = contextual
        return questions

    def _generate_contextual_question(self, placeholder):
        placeholder = placeholder.strip().upper()
        context_map = {
            "DATE": "What is the effective date of the contract?",
            "DISCLOSING_PARTY_NAME": "Who is the disclosing party?",
            "RECEIVING_PARTY_NAME": "Who will receive the information?",
            "CONFIDENTIAL_INFO_DESCRIPTION": "What information is considered confidential?",
            "DURATION": "What is the duration of this contract?"
        }
        return context_map.get(placeholder, "FALLBACK")

    def _postprocess_question(self, q):
        q = q.strip()
        if q.endswith("?"):
            return q
        if not q:
            return "Please provide a value."
        return q[0].capitalize() + q[1:] + "?"

    def ask_next_question(self, user_id="default_user", answer=None) -> str:
        session = CONTRACT_SESSION.get(user_id, {})
        if not session:
            # Init session
            self.extract_placeholders()
            questions = self.generate_questions()
            session = {
                "questions": questions,
                "fields": list(questions.keys()),
                "answers": {},
                "current_index": 0,
                "template": self.template_path
            }
            CONTRACT_SESSION[user_id] = session
            return f"📄 Starting contract creation.\n{questions[session['fields'][0]]}"

        if answer is not None:
            current_field = session["fields"][session["current_index"]]
            session["answers"][current_field] = answer
            session["current_index"] += 1

        if session["current_index"] >= len(session["fields"]):
            # All answered → fill contract
            self.responses = session["answers"]
            CONTRACT_SESSION.pop(user_id, None)

            self.fill_placeholders(self.responses)
            self.export_responses_pdf()
            return (
                f"✅ Contract completed!\n"
                f"📄 DOCX: `{self.docx_output}`\n📑 PDF: `{self.pdf_output}`"
            )

        # Ask next question
        next_field = session["fields"][session["current_index"]]
        question = session["questions"][next_field]
        CONTRACT_SESSION[user_id] = session
        return question

    def fill_placeholders(self, responses: dict):
        try:
            self.responses = responses
            doc = Document(self.template_path)

            for para in doc.paragraphs:
                for field, answer in self.responses.items():
                    if f"[{field}]" in para.text:
                        para.text = para.text.replace(f"[{field}]", str(answer))

            doc.save(self.docx_output)
            return self.docx_output
        except Exception as e:
            logging.error(f"Error filling document: {e}")
            raise

    def fill_single_contract(self, data: dict) -> str:
        """
        Public method used in smart agent for single contract generation.
        """
        self.responses = data
        self.fill_placeholders(data)
        self.export_responses_pdf()
        return self.docx_output

    def export_responses_pdf(self) -> str:
        try:
            c = canvas.Canvas(self.pdf_output, pagesize=letter)
            c.setFont("Helvetica-Bold", 14)
            c.drawString(100, 750, "User Responses for Contract")
            c.setFont("Helvetica", 12)

            y = 730
            for field, response in self.responses.items():
                if y < 50:
                    c.showPage()
                    c.setFont("Helvetica", 12)
                    y = 750
                c.drawString(100, y, f"{field}: {response}")
                y -= 20

            c.save()
            return self.pdf_output
        except Exception as e:
            logging.error(f"Failed to export PDF: {e}")
            raise

    def prepare_excel_template(self) -> str:
        try:
            df = pd.DataFrame(columns=self.placeholders)
            path = os.path.join("outputs", f"{os.path.basename(self.output_base)}_template.xlsx")
            df.to_excel(path, index=False)
            return path
        except Exception as e:
            logging.error(f"Error creating Excel template: {e}")
            raise

    def generate_bulk_contracts(self, filled_excel_path: str) -> list:
        try:
            df = pd.read_excel(filled_excel_path)
            paths = []
            for idx, row in df.iterrows():
                doc = Document(self.template_path)
                for para in doc.paragraphs:
                    for field in df.columns:
                        placeholder = f"[{field}]"
                        if placeholder in para.text:
                            para.text = para.text.replace(placeholder, str(row[field]))
                output_path = os.path.join("outputs", f"{os.path.basename(self.output_base)}_{idx + 1}.docx")
                doc.save(output_path)
                paths.append(output_path)
            return paths
        except Exception as e:
            logging.error(f"Bulk generation failed: {e}")
            raise
